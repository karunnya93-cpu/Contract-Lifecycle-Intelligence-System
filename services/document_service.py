from __future__ import annotations

from pathlib import Path


class DocumentExtractionError(RuntimeError):
    pass


class DocumentTextExtractor:
    @staticmethod
    def extract_text(file_path: Path) -> str:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return DocumentTextExtractor.extract_pdf_text(file_path)
        if suffix == ".docx":
            return DocumentTextExtractor.extract_docx_text(file_path)
        raise DocumentExtractionError("Only PDF and DOCX contracts are supported.")

    @staticmethod
    def extract_pdf_text(file_path: Path) -> str:
        text = DocumentTextExtractor._extract_with_pdfplumber(file_path)
        if not text:
            text = DocumentTextExtractor._extract_with_pypdf2(file_path)

        cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        if not cleaned:
            raise DocumentExtractionError("No readable text could be extracted from this PDF.")
        return cleaned

    @staticmethod
    def extract_docx_text(file_path: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentExtractionError(
                "DOCX support is unavailable. Install the python-docx package."
            ) from exc

        try:
            document = Document(str(file_path))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            cleaned = "\n".join(line.strip() for line in paragraphs + table_cells if line.strip())
        except Exception as exc:
            raise DocumentExtractionError(f"DOCX text extraction failed: {exc}") from exc

        if not cleaned:
            raise DocumentExtractionError("No readable text could be extracted from this DOCX file.")
        return cleaned

    @staticmethod
    def _extract_with_pdfplumber(file_path: Path) -> str:
        try:
            import pdfplumber
        except ImportError:
            return ""

        try:
            with pdfplumber.open(file_path) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages]
            return "\n".join(pages)
        except Exception as exc:
            raise DocumentExtractionError(f"PDF text extraction failed: {exc}") from exc

    @staticmethod
    def _extract_with_pypdf2(file_path: Path) -> str:
        try:
            from PyPDF2 import PdfReader
        except ImportError as exc:
            raise DocumentExtractionError(
                "Install pdfplumber or PyPDF2 to extract text from uploaded contracts."
            ) from exc

        try:
            reader = PdfReader(str(file_path))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages)
        except Exception as exc:
            raise DocumentExtractionError(f"Fallback PDF text extraction failed: {exc}") from exc
